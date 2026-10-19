"""HRMS+ESS v3.1 tests: reports & Excel export hub, retail-fixture registers,
SVG menu icons, PJAX app-like navigation, separation fields.

Uses an isolated temp SQLite DB and test fixtures only (no real employee data).
Run: ESS_SECRET_KEY=testsecret12345 ESS_DATABASE_URL=<redacted> \\
     python test_v31.py
"""
import io
import json
import os
import re
import sys
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from app import create_app, db  # noqa: E402
from app.models import (User, Company, PayrollRun, PayrollLine, PayrollInput,
                        LeaveBalance, LeaveRequest, AttendanceRecord,
                        SeparationRecord, Certification, Project, Deployment,
                        FacilityAssignment, PPEIssue, HSEIncident,
                        ContractorWorker, ReportTemplate, AuditLog)  # noqa: E402
import config as app_config  # noqa: E402

app = create_app()
app.config["WTF_CSRF_ENABLED"] = True
from app import limiter  # noqa: E402
limiter.enabled = False  # many logins in this suite; rate limiting is live-tested separately
client = app.test_client()

passed, failed = [], []


def check(name, cond, extra=""):
    (passed if cond else failed).append(name)
    print(("PASS " if cond else "FAIL ") + name + (f" [{extra}]" if extra else ""))


def csrf_token(resp):
    m = re.search(rb'name="csrf_token" value="([^"]+)"', resp.data)
    return m.group(1).decode() if m else ""


def get_token(path):
    return csrf_token(client.get(path))


def login(uid, pw):
    r = client.get("/login")
    t = csrf_token(r)
    return client.post("/login", data={"username": uid, "password": pw,
                                       "csrf_token": t},
                       follow_redirects=False)


def logout():
    t = csrf_token(client.get("/change-password"))
    return client.post("/logout", data={"csrf_token": t},
                       follow_redirects=False)


def mk(uid, name, role, dept, **kw):
    if not User.query.get(uid):
        kw.setdefault("joining_date", date(2020, 1, 15))
        u = User(id=uid, name=name, role=role, department=dept,
                 employment_status="active", is_active=True, **kw)
        u.set_password("pw-" + uid)
        u.must_change_password = False
        db.session.add(u)
    return User.query.get(uid)


def seed_master():
    for code, name in app_config.DEFAULT_COMPANIES:
        if not Company.query.filter_by(code=code).first():
            db.session.add(Company(code=code, name=name))
    db.session.commit()


def xlsx_sheets(resp_data):
    from openpyxl import load_workbook
    wb = load_workbook(filename=io.BytesIO(resp_data))
    ws = wb.active
    header = None
    for row in ws.iter_rows(values_only=True):
        vals = [str(v) for v in row if v]
        if "Employee ID" in vals or "Metric" in vals:
            header = vals
            break
    return ws, header


today = date.today()
PERIOD = "2026-10"

with app.app_context():
    seed_master()
    admin = mk("V31-ADM", "V31 Admin", "admin", "Human Resources",
               basic_salary=9000, monthly_salary=14000, iban="AE071")
    mgr = mk("V31-MGR", "V31 Manager", "manager", "Production",
             basic_salary=6000, monthly_salary=9000)
    emp = mk("V31-E1", "V31 Employee One", "employee", "Production",
             designation="Welder", basic_salary=3000, monthly_salary=4500,
             reporting_manager="V31-MGR")

    # ---- payroll run fixture for add/ded + OT summary ----
    run = PayrollRun(period=PERIOD, status="draft", created_by="V31-ADM")
    db.session.add(run)
    db.session.flush()
    db.session.add(PayrollLine(
        run_id=run.id, user_id="V31-E1", payable_days=30, lop_days=0,
        basic_earned=3000,
        earnings_json=json.dumps([
            {"code": "BASIC", "name": "Basic Salary", "amount": 3000},
            {"code": "HRA", "name": "Housing Allowance", "amount": 1000},
            {"code": "BONUS", "name": "Bonus", "amount": 200,
             "kind": "addition", "is_arrears": False}]),
        deductions_json=json.dumps([
            {"code": "LOAN", "name": "Loan Recovery", "amount": 500}]),
        ot_json=json.dumps([
            {"category": "normal", "hours": 10, "multiplier": 1.25,
             "source": "basic", "amount": 187.5}]),
        gross=4387.5, total_deductions=500, net=3887.5))
    db.session.add(PayrollInput(user_id="V31-E1", period=PERIOD, kind="addition",
                               amount=200, is_arrears=False,
                               label="Bonus", created_by="V31-ADM"))
    # ---- absence fixture: 20 present + 2 absent -> 9.1% ----
    for i in range(20):
        db.session.add(AttendanceRecord(user_id="V31-E1",
                                        date=date(2026, 10, 1) + timedelta(days=i),
                                        status="present"))
    db.session.add(AttendanceRecord(user_id="V31-E1", date=date(2026, 10, 21),
                                    status="absent"))
    db.session.add(AttendanceRecord(user_id="V31-E1", date=date(2026, 10, 22),
                                    status="absent"))
    # ---- leave ledger fixture: entitled 30, used 5, pending 2 -> remaining 23 ----
    db.session.add(LeaveBalance(user_id="V31-E1", leave_type="annual",
                                entitled=30, used=5))
    db.session.add(LeaveRequest(user_id="V31-E1", leave_type="annual",
                                start_date=date(2026, 11, 1),
                                end_date=date(2026, 11, 2), days=2,
                                status="pending"))
    # ---- cert expiring within 60 days for the compliance test ----
    db.session.add(Certification(user_id="V31-E1", name="Forklift Operation",
                                 issuer="Test Academy",
                                 issue_date=today - timedelta(days=700),
                                 expiry_date=today + timedelta(days=30),
                                 created_by="V31-ADM"))
    db.session.commit()

# ================= reports hub =================
r = login("V31-ADM", "pw-V31-ADM")
check("admin login ok", r.status_code in (301, 302, 303))

r = client.get("/hrms/reports")
check("reports hub 200 for admin", r.status_code == 200)
check("hub lists master export", b"Master Data Export" in r.data)
check("hub lists turnover", b"Turnover Report" in r.data)

logout()
login("V31-E1", "pw-V31-E1")
r = client.get("/hrms/reports")
check("reports hub 403 for plain employee", r.status_code == 403)
logout()
login("V31-ADM", "pw-V31-ADM")

# ================= (a) master data export =================
r = client.get("/hrms/reports/master?format=xlsx")
check("master export 200 xlsx", r.status_code == 200 and
      "spreadsheetml" in r.headers.get("Content-Type", ""))
ws, header = xlsx_sheets(r.data)
check("master xlsx header found", header is not None)
check("master xlsx has Employee ID header", "Employee ID" in (header or []))
check("master xlsx has salary column for admin", "Basic (AED)" in (header or []))
check("master xlsx has IBAN column for admin", "IBAN" in (header or []))
found_admin = any("V31-ADM" in [str(v) for v in row if v]
                  for row in ws.iter_rows(values_only=True))
check("master xlsx contains a known employee row", found_admin)

r = client.get("/hrms/reports/master")
check("master html view 200", r.status_code == 200 and b"Master Data Export" in r.data)

# ================= (b) additions & deductions =================
r = client.get("/hrms/reports/additions-deductions?period=" + PERIOD)
check("add/ded html 200", r.status_code == 200)
check("add/ded shows employee net matching the run line",
      b"3887.50" in r.data)
check("add/ded shows draft status note", b"draft" in r.data)
check("add/ded shows of-which inputs breakout", b"of which: Inputs" in r.data)

logout()
login("V31-MGR", "pw-V31-MGR")  # manager WITHOUT salary permission
r = client.get("/hrms/reports/additions-deductions?period=" + PERIOD)
check("add/ded 403 for manager without salary permission",
      r.status_code == 403)
r = client.get("/hrms/reports/ot-summary?period=" + PERIOD)
check("ot summary 403 for manager without salary permission",
      r.status_code == 403)
logout()
login("V31-ADM", "pw-V31-ADM")

r = client.get("/hrms/reports/ot-summary?period=" + PERIOD)
check("ot summary html 200", r.status_code == 200)
check("ot summary shows normal OT hours", b"10.0" in r.data)

# ================= (c) absence report =================
r = client.get("/hrms/reports/absence?month=2026-10")
check("absence html 200", r.status_code == 200)
check("absence rate math: 2 absent of 22 -> 9.1%",
      b"9.1%" in r.data)
check("absence shows bradford column with footnote",
      b"Bradford Factor" in r.data and b"optional" in r.data)

# ================= (d) leave report =================
r = client.get("/hrms/reports/leave")
check("leave ledger 200", r.status_code == 200)
check("leave ledger remaining = 30 - 5 - 2 = 23.00", b"23.00" in r.data)
r = client.get("/hrms/reports/leave?view=requests")
check("leave requests view 200", r.status_code == 200 and b"Pending" in r.data)

# ================= (e) separation CRUD + report =================
t = get_token("/hrms/separations")
r = client.get("/hrms/separations")
check("separations page 200 for admin", r.status_code == 200)
r = client.post("/hrms/separations", data={
    "csrf_token": t, "employee_code": "AA991",
    "employee_name": "Test Separated", "department": "Production",
    "designation": "Welder", "joining_date": "2021-03-01",
    "separation_date": "2024-02-15", "separation_type": "Resignation",
    "reason": "Test", "notice_days": "30",
    "final_settlement_status": "Settled", "settlement_amount": "8500"})
check("separation record created (redirect)", r.status_code in (301, 302, 303))
r = client.get("/hrms/separations")
check("separations list shows the record", b"Test Separated" in r.data)
r = client.get("/hrms/reports/separations")
check("separation report contains the record", b"Test Separated" in r.data)
r = client.get("/hrms/reports/separations?format=xlsx")
check("separation export 200 xlsx", r.status_code == 200)
ws2, _ = xlsx_sheets(r.data)
found_sep = any("Test Separated" in [str(v) for v in row if v]
                for row in ws2.iter_rows(values_only=True))
check("separation xlsx contains the record", found_sep)

logout()
login("V31-E1", "pw-V31-E1")
r = client.get("/hrms/separations")
check("separations 403 for plain employee", r.status_code == 403)
logout()
login("V31-ADM", "pw-V31-ADM")

# ================= (f) turnover =================
with app.app_context():
    for i in range(10):
        mk(f"V31-T{i:02d}", f"V31 Turnover {i:02d}", "employee",
           "Production", joining_date=date(2025, 6, 1))
    mk("V31-J1", "V31 Joiner", "employee", "Production",
       joining_date=date(2026, 3, 1))
    db.session.add(SeparationRecord(
        employee_code="AA981", employee_name="Turnover A",
        department="Production", joining_date=date(2022, 1, 1),
        separation_date=date(2026, 4, 10), separation_type="Resignation",
        final_settlement_status="Settled", created_by="V31-ADM"))
    db.session.add(SeparationRecord(
        employee_code="AA982", employee_name="Turnover B",
        department="Production", joining_date=date(2023, 1, 1),
        separation_date=date(2026, 9, 5), separation_type="Termination",
        final_settlement_status="Pending", created_by="V31-ADM"))
    db.session.commit()
# year 2026: separations=2, joiners=1 (V31-J1), end_hc=14 (all users),
# start_hc = 14 - 1 + 2 = 15, avg = 14.5, rate = 2/14.5*100 = 13.8%
r = client.get("/hrms/reports/turnover?mode=year&year=2026")
check("turnover 200", r.status_code == 200)
check("turnover rate matches hand calc (13.8%)", b"13.8%" in r.data)
check("turnover page shows the formula", b"Turnover %" in r.data and
      b"Average headcount" in r.data)
check("turnover voluntary/involuntary split shown",
      b"Voluntary" in r.data and b"Involuntary" in r.data)
check("turnover per-department table shown", b"Production" in r.data)

# ================= (h) loan balances =================
r = client.get("/hrms/reports/loans")
check("loan balances 200", r.status_code == 200)

# ================= every export writes an audit row =================
exports = [
    ("master", "/hrms/reports/master?format=xlsx"),
    ("addded", "/hrms/reports/additions-deductions?format=xlsx&period=" + PERIOD),
    ("absence", "/hrms/reports/absence?format=xlsx&month=2026-10"),
    ("leave", "/hrms/reports/leave?format=xlsx"),
    ("separations", "/hrms/reports/separations?format=xlsx"),
    ("turnover", "/hrms/reports/turnover?format=xlsx&mode=year&year=2026"),
    ("otsummary", "/hrms/reports/ot-summary?format=xlsx&period=" + PERIOD),
    ("loans", "/hrms/reports/loans?format=xlsx"),
]
with app.app_context():
    before = {k: AuditLog.query.filter_by(action="report_export",
                                           entity_id=k).count()
              for k, _ in exports}
for key, url in exports:
    r = client.get(url)
    check(f"export {key} returns 200 xlsx",
          r.status_code == 200 and "spreadsheetml" in
          r.headers.get("Content-Type", ""))
with app.app_context():
    for key, _ in exports:
        n = AuditLog.query.filter_by(action="report_export",
                                     entity_id=key).count()
        check(f"export {key} wrote an audit row", n > before[key])
        row = AuditLog.query.filter_by(action="report_export",
                                       entity_id=key)\
            .order_by(AuditLog.id.desc()).first()
        check(f"audit row for {key} names actor + report + filters",
              row is not None and "V31-ADM" in (row.actor_id or "") and
              "filters:" in (row.detail or ""))

# ================= (i) saved report templates =================
t = get_token("/hrms/reports/absence?month=2026-10")
r = client.post("/hrms/reports/templates/save", data={
    "csrf_token": t, "report_key": "absence", "name": "Oct absence",
    "params": '{"month": "2026-10"}'}, follow_redirects=False)
check("template saved (redirect)", r.status_code in (301, 302, 303))
r = client.get("/hrms/reports/absence?month=2026-10")
check("saved template appears on the report page", b"Oct absence" in r.data)

# ================= registers =================
t = get_token("/hrms/certifications")
r = client.post("/hrms/certifications", data={
    "csrf_token": t, "user_id": "V31-E1", "name": "Welding Trade",
    "issuer": "Test Institute", "issue_date": "2024-01-01",
    "expiry_date": "2027-01-01"}, follow_redirects=False)
check("certification created by admin", r.status_code in (301, 302, 303))
r = client.get("/hrms/certifications")
check("certifications list 200", r.status_code == 200)
check("certifications list shows the cert", b"Welding Trade" in r.data)

r = client.post("/hrms/projects", data={
    "csrf_token": t, "name": "Test Fit-Out", "client": "Test Client",
    "location": "Dubai", "status": "active"}, follow_redirects=False)
check("project created by admin", r.status_code in (301, 302, 303))
with app.app_context():
    pid = Project.query.filter_by(name="Test Fit-Out").first().id
r = client.get(f"/hrms/projects/{pid}")
check("project detail 200", r.status_code == 200)
r = client.post(f"/hrms/projects/{pid}", data={
    "csrf_token": t, "user_id": "V31-E1",
    "from_date": "2026-09-01", "site_allowance": "250"},
    follow_redirects=False)
check("deployment created by admin", r.status_code in (301, 302, 303))
r = client.get(f"/hrms/projects/{pid}")
check("project detail shows total site allowance",
      b"250.00" in r.data)

r = client.post("/hrms/facilities", data={
    "csrf_token": t, "user_id": "V31-E1", "kind": "accommodation",
    "location_or_route": "Camp C", "monthly_cost": "250",
    "from_date": "2026-01-01"}, follow_redirects=False)
check("facility assignment created", r.status_code in (301, 302, 303))
r = client.post("/hrms/ppe", data={
    "csrf_token": t, "user_id": "V31-E1", "item": "Safety Helmet",
    "size": "M", "qty": "1", "issued_date": "2026-09-01"},
    follow_redirects=False)
check("ppe issue created", r.status_code in (301, 302, 303))

logout()
login("V31-MGR", "pw-V31-MGR")
t = get_token("/hrms/hse")
r = client.post("/hrms/hse", data={
    "csrf_token": t, "date": "2026-10-10", "location": "Welding Bay",
    "type": "Near Miss", "severity": "Low",
    "description": "Spark near storage"}, follow_redirects=False)
check("hse incident created by manager", r.status_code in (301, 302, 303))
r = client.get("/hrms/certifications")
check("certifications readable by manager", r.status_code == 200)
r = client.post("/hrms/certifications", data={
    "csrf_token": t, "user_id": "V31-E1", "name": "X"},
    follow_redirects=False)
check("certification create 403 for manager", r.status_code == 403)
logout()
login("V31-E1", "pw-V31-E1")
r = client.get("/hrms/certifications")
check("certifications 403 for plain employee", r.status_code == 403)
r = client.get("/hrms/hse")
check("hse 403 for plain employee", r.status_code == 403)
r = client.get("/hrms/projects")
check("projects 403 for plain employee", r.status_code == 403)
r = client.get("/profile")
check("employee profile shows own certification", b"Forklift Operation" in r.data)
logout()
login("V31-ADM", "pw-V31-ADM")

t = get_token("/hrms/contractors")
r = client.post("/hrms/contractors", data={
    "csrf_token": t, "name": "Test Worker", "company": "Test Manpower",
    "trade": "Welder", "site": "UAQ", "monthly_rate": "1800",
    "status": "active"}, follow_redirects=False)
check("contractor created", r.status_code in (301, 302, 303))
r = client.get("/hrms/contractors")
check("contractor billing summary shown",
      b"1,800.00" in r.data and b"Test Manpower" in r.data)

# ================= compliance shows expiring cert =================
r = client.get("/hrms/compliance")
check("compliance 200", r.status_code == 200)
check("compliance lists expiring certification",
      b"Certification Expiries" in r.data and b"Forklift Operation" in r.data)

# ================= SVG icons (no emojis) =================
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "app", "templates", "base.html"),
          encoding="utf-8") as fh:
    base_html = fh.read()
OLD_EMOJI = [127891, 127970, 128100, 128101, 128172, 128176, 128178, 128179,
             128196, 128197, 128202, 128203, 128221, 128241, 128336, 128451,
             8681, 8962, 9200, 9733, 9776, 9878, 9989, 9993, 9998]
emoji_found = [c for c in OLD_EMOJI if chr(c) in base_html]
check("base.html nav has no old emoji chars", not emoji_found,
      extra="found: " + ",".join(map(str, emoji_found)))
check("base.html has no old emoji entity refs",
      not any(f"&#{c}" in base_html for c in OLD_EMOJI))
check("base.html nav calls the svg icon macro", "icon(" in base_html)
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "app", "templates", "_icons.html"),
          encoding="utf-8") as fh:
    icons_html = fh.read()
check("icon macro file contains inline svg", "<svg" in icons_html)
check("icon macro file has no old emoji chars",
      not any(chr(c) in icons_html for c in OLD_EMOJI))
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "app", "templates", "dashboard.html"),
          encoding="utf-8") as fh:
    dash_html = fh.read()
check("dashboard tiles have no old emoji chars",
      not any(chr(c) in dash_html for c in OLD_EMOJI))
check("dashboard tiles call the svg icon macro", "icon(" in dash_html)
# and the proof that matters: the rendered page carries real <svg> icons
r = client.get("/")
check("rendered navigation contains inline svg icons",
      r.status_code == 200 and b"<svg" in r.data)
logout()
login("V31-ADM", "pw-V31-ADM")

# ================= PJAX navigation =================
check("base.html has the pjax container", 'id="pjax-main"' in base_html)
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "app", "static", "app.js"), encoding="utf-8") as fh:
    app_js = fh.read()
check("app.js contains pjax interception", "pjax" in app_js)
check("app.js honours data-no-pjax", "data-no-pjax" in app_js)
check("app.js handles back/forward (popstate)", "popstate" in app_js)
check("app.js updates history (pushState)", "pushState" in app_js)
check("app.js excludes downloads/exports", "/export" in app_js)
r = client.get("/")
check("normal page GET still returns full HTML 200",
      r.status_code == 200 and b"<html" in r.data and b'id="pjax-main"' in r.data)
r = client.get("/static/app.js")
check("static app.js served with long cache header",
      r.status_code == 200 and
      "max-age" in r.headers.get("Cache-Control", ""))
logout()
r = client.get("/login")
check("login page unaffected by pjax", r.status_code == 200 and
      b"pjax-main" not in r.data)

print(f"\n{len(passed)} passed, {len(failed)} failed")
sys.exit(1 if failed else 0)
