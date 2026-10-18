"""HRMS+ESS v3.1.1 tests: Gemini Argon review round.

- XLSX formula-injection sanitisation
- Export watermarking (properties + hidden Export Info sheet)
- Payroll variance report, Leave & EOSB liability report, WPS pre-flight check
- Training matrix, Asset & Tool Custody register + separation warnings

Uses an isolated temp SQLite DB and test fixtures only (no real employee data).
Run: ESS_SECRET_KEY=testsecret12345 ESS_DATABASE_URL=<redacted> \\
     python test_v311.py
"""
import io
import json
import os
import re
import sys
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from app import create_app, db  # noqa: E402
from app.models import (User, Company, PayrollRun, PayrollLine, LeaveBalance,
                        SeparationRecord, Certification, AssetCustody,
                        AuditLog)  # noqa: E402
from app.uae import gratuity as gratuity_calc  # noqa: E402
import config as app_config  # noqa: E402

app = create_app()
app.config["WTF_CSRF_ENABLED"] = True
from app import limiter  # noqa: E402
limiter.enabled = False
client = app.test_client()

passed, failed = [], []


def check(name, cond, extra=""):
    (passed if cond else failed).append(name)
    print(("PASS " if cond else "FAIL ") + name + (f" [{extra}]" if extra else ""))


def csrf_token(resp):
    m = re.search(rb'name="csrf_token" value="([^"]+)"', resp.data)
    return m.group(1).decode() if m else ""


def get_token(path):
    return csrf_token(client.get(path))


def login(uid, pw):
    r = client.get("/login")
    t = csrf_token(r)
    return client.post("/login", data={"username": uid, "password": pw,
                                       "csrf_token": t},
                       follow_redirects=False)


def logout():
    t = csrf_token(client.get("/change-password"))
    return client.post("/logout", data={"csrf_token": t},
                       follow_redirects=False)


def mk(uid, name, role, dept, **kw):
    if not User.query.get(uid):
        kw.setdefault("joining_date", date(2020, 1, 15))
        u = User(id=uid, name=name, role=role, department=dept,
                 employment_status="active", is_active=True, **kw)
        u.set_password("pw-" + uid)
        u.must_change_password = False
        db.session.add(u)
    return User.query.get(uid)


def load_xlsx(data):
    from openpyxl import load_workbook
    return load_workbook(filename=io.BytesIO(data))


today = date.today()

with app.app_context():
    for code, name in app_config.DEFAULT_COMPANIES:
        if not Company.query.filter_by(code=code).first():
            db.session.add(Company(code=code, name=name))
    db.session.commit()
    comp = Company.query.filter_by(code=app_config.DEFAULT_COMPANIES[0][0]).first()

    admin = mk("V11-ADM", "V11 Admin", "admin", "Human Resources",
               basic_salary=12000, monthly_salary=18000,
               emirates_id_no="784-1234-1234567-1",
               iban="AE" + "0" * 21, bank_name="Demo Bank",
               visa_company_id=comp.id)
    mgr = mk("V11-MGR", "V11 Manager", "manager", "Production",
             basic_salary=6000, monthly_salary=9000)
    e1 = mk("V11-E1", "V11 Employee One", "employee", "Production",
            basic_salary=3000, monthly_salary=5000,
            emirates_id_no="784-9999-8888888-8",
            iban="AE" + "1" * 21, bank_name="Demo Bank",
            visa_company_id=comp.id)
    e2 = mk("V11-E2", "V11 Employee Two", "employee", "Production",
            basic_salary=2500, monthly_salary=4200,
            emirates_id_no=None, iban="123",
            bank_name="", visa_company_id=None)
    e3 = mk("V11-E3", "V11 Employee Three", "employee", "Production",
            basic_salary=2200, monthly_salary=3800)
    e4 = mk("V11-E4", "V11 Employee Four", "employee", "Production",
            basic_salary=2000, monthly_salary=3400)

    def line(run, uid, net, gross=None, ded=0.0):
        db.session.add(PayrollLine(
            run_id=run.id, user_id=uid, payable_days=30, lop_days=0,
            basic_earned=3000,
            earnings_json=json.dumps([{"code": "BASIC", "name": "Basic",
                                       "amount": 3000}]),
            deductions_json=json.dumps([]),
            ot_json=json.dumps([]),
            gross=gross if gross is not None else net,
            total_deductions=ded, net=net))

    prev = PayrollRun(period="2026-09", status="approved", created_by="V11-ADM")
    cur = PayrollRun(period="2026-10", status="locked", created_by="V11-ADM")
    db.session.add_all([prev, cur])
    db.session.flush()
    line(prev, "V11-E1", 4000)          # +12.5% in Oct -> exception
    line(prev, "V11-E2", 3000)          # unchanged
    line(prev, "V11-E4", 2500)          # leaver
    line(cur, "V11-E1", 4500)
    line(cur, "V11-E2", 3000)
    line(cur, "V11-E3", 2000)           # new joiner
    # wps pre-flight run: one clean, one bad
    wps = PayrollRun(period="2026-11", status="locked", created_by="V11-ADM")
    db.session.add(wps)
    db.session.flush()
    line(wps, "V11-E1", 4500, gross=5000, ded=500)
    line(wps, "V11-E2", 2000, gross=5000, ded=3000)   # bad: IBAN/EID/bank/company + ratios
    solo = PayrollRun(period="2027-01", status="draft", created_by="V11-ADM")
    db.session.add(solo)
    db.session.flush()
    line(solo, "V11-E1", 4500)
    db.session.commit()

login("V11-ADM", "pw-V11-ADM")

# ================= 1. formula-injection sanitisation =================
with app.app_context():
    payload = "=2+5+cmd|'/c calc'!A0"
    u = User.query.get("V11-E1")
    u.name = payload
    db.session.commit()
r = client.get("/hrms/reports/master?format=xlsx")
wb = load_xlsx(r.data)
ws = wb.active
found = None
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value, str) and "cmd" in c.value:
            found = c
check("payload cell found in export", found is not None)
check("payload cell neutralised with quote prefix",
      found is not None and found.value == "'" + payload)
check("payload cell is a string, not a formula",
      found is not None and found.data_type != "f")
check("xlsx sanitiser leaves normal text alone",
      any(c.value == "V11-E3" for row in ws.iter_rows() for c in row))

# ================= 2. export watermarking =================
check("workbook creator names the exporting user",
      wb.properties.creator == "V11-ADM (V11 Admin)")
check("workbook title names the report",
      wb.properties.title == "Master Data Export")
info = wb["Export Info"]
check("hidden Export Info sheet exists", info is not None)
check("Export Info sheet is hidden",
      info.sheet_state in ("hidden", "veryHidden"))
vals = {r[0]: r[1] for r in info.iter_rows(values_only=True)}
check("Export Info names the exporter", "V11-ADM" in str(vals.get("Exported by")))
check("Export Info carries a UTC timestamp",
      "T" in str(vals.get("Exported at (UTC)")) and str(vals.get("Exported at (UTC)")).endswith("Z"))
check("Export Info carries the filters", "department=" in str(vals.get("Filters used")))

# ================= 3. variance report =================
r = client.get("/hrms/reports/variance?period=2026-10")
check("variance html 200", r.status_code == 200)
check("variance shows changed employee with +12.5%",
      b"12.5%" in r.data and b"4500.00" in r.data)
check("variance shows joiner bucket", b"New Joiner" in r.data)
check("variance shows leaver bucket", b"Leaver" in r.data)
check("variance shows bridge totals", b"9500.00" in r.data)
check("variance exceptions section lists the >10% mover",
      b"Exceptions" in r.data and r.data.count(b"V11-E1") >= 2)
r = client.get("/hrms/reports/variance?period=2026-10&format=xlsx")
check("variance xlsx 200", r.status_code == 200)
wb = load_xlsx(r.data)
check("variance xlsx has Exceptions sheet", "Exceptions" in wb.sheetnames)
exc = wb["Exceptions"]
exc_vals = [row[0] for row in exc.iter_rows(values_only=True)]
check("exceptions sheet contains the exception row", "V11-E1" in exc_vals)
check("exceptions sheet excludes unchanged, joiner and leaver",
      "V11-E2" not in exc_vals and "V11-E3" not in exc_vals
      and "V11-E4" not in exc_vals)
r = client.get("/hrms/reports/variance?period=2027-01")
check("variance with no previous run shows a clear message",
      r.status_code == 200 and b"No previous payroll run" in r.data)

# ================= 4. liability report =================
with app.app_context():
    db.session.add(LeaveBalance(user_id="V11-E1", leave_type="annual",
                                entitled=30, used=5))
    db.session.commit()
    expected_grat = gratuity_calc(date(2020, 1, 15), today, 3000, "employer")["amount"]
    expected_leave = round(25.0 * (3000 / 30.0), 2)
r = client.get("/hrms/reports/liability")
check("liability html 200", r.status_code == 200)
check("liability gratuity matches independent calc",
      f"{expected_grat:.2f}".encode() in r.data)
check("liability leave value matches hand calc (25 days x 100)",
      f"{expected_leave:.2f}".encode() in r.data)
r = client.get("/hrms/reports/liability?format=xlsx")
check("liability xlsx 200", r.status_code == 200)
wb = load_xlsx(r.data)
check("liability xlsx watermarked", wb["Export Info"].sheet_state in
      ("hidden", "veryHidden"))

# ================= 5. wps pre-flight =================
r = client.get("/hrms/reports/wps-preflight?period=2026-11")
check("wps pre-flight 200", r.status_code == 200)
check("pre-flight flags invalid IBAN", b"Invalid IBAN" in r.data)
check("pre-flight flags missing Emirates ID", b"Missing Emirates ID number" in r.data)
check("pre-flight flags missing bank", b"Missing bank name" in r.data)
check("pre-flight flags missing visa company", b"Missing visa company" in r.data)
check("pre-flight flags net below 80% of gross", b"Net pay below 80% of gross" in r.data)
check("pre-flight flags deductions above 50%", b"Total deductions above 50% of gross" in r.data)
check("pre-flight summary counts 1 OK and 1 flagged",
      b"1 OK" in r.data and b"1 flagged" in r.data)
check("pre-flight does not flag the clean employee",
      b"V11-E1" not in r.data)
check("pre-flight flags only the bad employee",
      r.data.count(b"V11-E2") >= 1)
r = client.get("/hrms/reports/wps-preflight?period=2026-11&format=xlsx")
check("pre-flight xlsx 200", r.status_code == 200)
wb = load_xlsx(r.data)
ids = [row[0] for row in wb.active.iter_rows(values_only=True)]
check("pre-flight xlsx lists only flagged employees",
      "V11-E2" in ids and "V11-E1" not in ids)

# ================= salary-gate on new money reports =================
logout()
login("V11-MGR", "pw-V11-MGR")  # no salary permission
for path, name in [("/hrms/reports/variance?period=2026-10", "variance"),
                   ("/hrms/reports/liability", "liability"),
                   ("/hrms/reports/wps-preflight?period=2026-11", "wps pre-flight")]:
    r = client.get(path)
    check(f"{name} 403 for manager without salary permission",
          r.status_code == 403)
logout()
login("V11-ADM", "pw-V11-ADM")

# ================= 6. training matrix =================
with app.app_context():
    db.session.add(Certification(user_id="V11-E1", name="Fire Safety",
                                 issuer="Test Academy",
                                 issue_date=today - timedelta(days=300),
                                 expiry_date=today + timedelta(days=30),
                                 created_by="V11-ADM"))
    db.session.add(Certification(user_id="V11-E2", name="Fire Safety",
                                 issuer="Test Academy",
                                 issue_date=today - timedelta(days=700),
                                 expiry_date=today - timedelta(days=10),
                                 created_by="V11-ADM"))
    db.session.commit()
r = client.get("/hrms/training-matrix")
check("training matrix 200", r.status_code == 200)
check("training matrix shows cert column", b"Fire Safety" in r.data)
check("training matrix shows expiring cert date",
      (today + timedelta(days=30)).isoformat().encode() in r.data)
check("training matrix marks expired cells", b"mx-expired" in r.data)
check("training matrix marks expiring cells", b"mx-expiring" in r.data)
check("training matrix dashes for employees without certs",
      b"mx-none" in r.data)
r = client.get("/hrms/training-matrix?department=Production")
check("training matrix department filter works",
      r.status_code == 200 and b"V11-E1" in r.data)
r = client.post("/hrms/training-matrix")
check("training matrix is read-only (POST 405)", r.status_code == 405)

logout()
login("V11-MGR", "pw-V11-MGR")
r = client.get("/hrms/training-matrix")
check("training matrix readable by manager", r.status_code == 200)
logout()
login("V11-E1", "pw-V11-E1")
r = client.get("/hrms/training-matrix")
check("training matrix 403 for plain employee", r.status_code == 403)
logout()
login("V11-ADM", "pw-V11-ADM")

# ================= 7. asset custody =================
t = get_token("/hrms/assets")
r = client.get("/hrms/assets")
check("assets page 200 for admin", r.status_code == 200)
r = client.post("/hrms/assets", data={
    "csrf_token": t, "user_id": "V11-E1", "item_name": "Power Drill",
    "category": "Tool", "serial_no": "PD-1001",
    "replacement_value": "650", "issued_date": "2026-08-01"},
    follow_redirects=False)
check("asset created by admin (redirect)", r.status_code in (301, 302, 303))
r = client.get("/hrms/assets")
check("assets list shows the item", b"Power Drill" in r.data)
with app.app_context():
    aid = AssetCustody.query.filter_by(item_name="Power Drill").first().id
r = client.post(f"/hrms/assets/{aid}/return", data={"csrf_token": t},
                follow_redirects=False)
check("asset return works (redirect)", r.status_code in (301, 302, 303))
r = client.get("/hrms/assets")
check("asset shows returned after return", b"Returned" in r.data)

# separation warning: E3 has an outstanding asset; record a separation
t = get_token("/hrms/separations")
r = client.post("/hrms/assets", data={
    "csrf_token": t, "user_id": "V11-E3", "item_name": "Angle Grinder",
    "category": "Tool", "serial_no": "AG-2002",
    "replacement_value": "400", "issued_date": "2026-07-01"},
    follow_redirects=False)
r = client.post("/hrms/separations", data={
    "csrf_token": t, "employee_code": "V11-E3",
    "employee_name": "V11 Employee Three", "department": "Production",
    "separation_date": "2026-10-15", "separation_type": "Resignation"},
    follow_redirects=False)
check("separation record for asset holder created",
      r.status_code in (301, 302, 303))
r = client.get("/hrms/separations")
check("separation register warns about unreturned assets",
      b"1 unreturned" in r.data)
r = client.get("/hrms/reports/separations")
check("separation report warns about unreturned assets",
      b"1 unreturned" in r.data)
logout()
login("V11-E3", "pw-V11-E3")
r = client.get("/profile")
check("employee profile shows own assets", b"Angle Grinder" in r.data)
logout()
login("V11-ADM", "pw-V11-ADM")

logout()
login("V11-MGR", "pw-V11-MGR")
r = client.get("/hrms/assets")
check("assets readable by manager", r.status_code == 200)
r = client.post("/hrms/assets", data={"csrf_token": get_token("/hrms/assets"),
                                       "user_id": "V11-E1",
                                       "item_name": "X"},
                follow_redirects=False)
check("asset create 403 for manager", r.status_code == 403)
logout()
login("V11-E1", "pw-V11-E1")
r = client.get("/hrms/assets")
check("assets 403 for plain employee", r.status_code == 403)
logout()
login("V11-ADM", "pw-V11-ADM")

# ================= new reports in hub + templates =================
r = client.get("/hrms/reports")
check("hub lists variance", b"Payroll Variance" in r.data)
check("hub lists liability", b"Leave &amp; EOSB Liability" in r.data)
check("hub lists wps pre-flight", b"WPS Pre-Flight" in r.data)
t = get_token("/hrms/reports/variance?period=2026-10")
r = client.post("/hrms/reports/templates/save", data={
    "csrf_token": t, "report_key": "variance", "name": "Oct variance",
    "params": '{"period": "2026-10"}'}, follow_redirects=False)
check("template saved for variance report", r.status_code in (301, 302, 303))
r = client.get("/hrms/reports/variance?period=2026-10")
check("variance page shows saved template", b"Oct variance" in r.data)

# ================= every new export wrote an audit row =================
with app.app_context():
    for key in ("variance", "liability", "wpspreflight"):
        n = AuditLog.query.filter_by(action="report_export",
                                     entity_id=key).count()
        check(f"export {key} wrote an audit row", n >= 1)

print(f"\n{len(passed)} passed, {len(failed)} failed")
sys.exit(1 if failed else 0)
